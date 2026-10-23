# -*- coding: utf-8 -*-
"""docx 格式规范校验 verify_docx_format.py —— 检查 Word 文档的排版规范。

设计原则（来自实战教训）：
- **按段落类型分组检查**：正文 / 标题 / 列表 / 题注 / 代码 / 其他 分别统计，
  不对"全部段落"套同一标准（否则题注、代码、封面会被误判）。
- **标准可配置**：默认用中文学术/技术文档的常见规范，可用 --config 覆盖；
  不做「>= 30 个段落缩进」这类臆断阈值。
- **只报「与规范不符」+「分布统计」**：缩进/行距本身不是硬性要求时只做统计，
  不判 FAIL，避免诱导无意义的"修复"。

用法：
  python verify_docx_format.py <doc.docx> [--config rules.json] [--json]
  python verify_docx_format.py doc.docx --profile academic   # 学术报告规范
  python verify_docx_format.py doc.docx --profile technical  # 技术文档规范（默认）

检查项：页面(A4/页边距)、标题层级样式、正文字体字号、首行缩进分布、行距分布、
        表格边框、图片尺寸、页眉页脚、代码块等宽字体。
"""
import sys, os, json, re, argparse
from collections import Counter, defaultdict
import docx
from docx.shared import Cm, Pt

# ── 默认规范（可按 profile 覆盖） ──────────────────────────────────────────
PROFILES = {
    'technical': {
        'page': {'width_cm': 21.0, 'height_cm': 29.7, 'tolerance_cm': 0.3},
        'body': {'fonts': ['宋体', '微软雅黑', 'SimSun', 'Microsoft YaHei'], 'size_pt': [10.5, 11, 12],
                 'indent_required': False, 'line_spacing_ok': [1.0, 1.15, 1.25, 1.5, 1.75, 2.0]},
        'heading': {'fonts': ['黑体', '微软雅黑', 'SimHei', 'Microsoft YaHei'], 'bold_required': True},
        'table': {'border_required': True},
        'image': {'max_width_margin_cm': 0.5},   # 图片宽度不得超过 页宽-左右边距-该余量
        'require_header': False, 'require_footer': False,
    },
    'academic': {
        'page': {'width_cm': 21.0, 'height_cm': 29.7, 'tolerance_cm': 0.3},
        'body': {'fonts': ['宋体', 'SimSun'], 'size_pt': [12], 'indent_required': True,
                 'indent_pt': [24, 21, 28],   # 约 2 字符
                 'line_spacing_ok': [1.25, 1.5]},
        'heading': {'fonts': ['黑体', 'SimHei', '宋体'], 'bold_required': True},
        'table': {'border_required': True},
        'image': {'max_width_margin_cm': 0.5},
        'require_header': False, 'require_footer': False,
    },
}

def load_rules(profile, config_path):
    rules = json.loads(json.dumps(PROFILES.get(profile, PROFILES['technical'])))
    if config_path:
        custom = json.load(open(config_path, encoding='utf-8'))
        for k, v in custom.items():
            if isinstance(v, dict) and k in rules:
                rules[k].update(v)
            else:
                rules[k] = v
    return rules

def style_group(p):
    """把段落归到检查分组。"""
    name = (p.style.name or '').lower()
    text = p.text.strip()
    if not text:
        return 'empty'
    if name.startswith('heading') or name in ('title', 'subtitle'):
        return 'heading'
    # 封面/大标题类：字号明显大于正文，或居中且短——不参与正文字体字号检查
    sizes = [r.font.size.pt for r in p.runs if r.text.strip() and r.font.size]
    if sizes and max(sizes) >= 16:
        return 'cover'
    if p.alignment is not None and len(text) <= 40 and 'CENTER' in str(p.alignment):
        return 'cover'
    if re.match(r'^(图|表)\s*\d+', text) and len(text) <= 40:
        return 'caption'
    if name.startswith('list'):
        return 'list'
    if all((r.font.name or '').lower().startswith(('consolas', 'courier')) for r in p.runs if r.text.strip()) and p.runs:
        return 'code'
    if text.startswith('>') or p.paragraph_format.left_indent is not None and not name.startswith('normal'):
        return 'other'
    return 'body'

def run_fonts(p):
    """返回段落内非空 run 的 (中文字体, 字号pt) 列表。"""
    out = []
    for r in p.runs:
        if not r.text.strip():
            continue
        east = None
        try:
            if r._element.rPr is not None and r._element.rPr.rFonts is not None:
                east = r._element.rPr.rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
        except Exception:
            pass
        size = r.font.size.pt if r.font.size else None
        out.append((east or r.font.name, size))
    return out

def table_has_borders(tbl):
    xml = tbl._tbl.xml
    return ('tblBorders' in xml) or (tbl.style is not None and 'Table Grid' in (tbl.style.name or ''))

def check(path, rules):
    d = docx.Document(path)
    issues = []      # 硬性问题
    notes = []       # 统计/建议（不判错）
    groups = defaultdict(list)
    for p in d.paragraphs:
        groups[style_group(p)].append(p)

    # 页面
    sec = d.sections[0]
    page = rules['page']
    w, h = sec.page_width.cm, sec.page_height.cm
    if abs(w - page['width_cm']) > page['tolerance_cm'] or abs(h - page['height_cm']) > page['tolerance_cm']:
        issues.append(f'页面尺寸 {w:.1f}x{h:.1f}cm，期望 A4 {page["width_cm"]}x{page["height_cm"]}cm')
    notes.append(f'页面：{w:.1f}x{h:.1f}cm；页边距 左{sec.left_margin.cm:.1f} 右{sec.right_margin.cm:.1f} 上{sec.top_margin.cm:.1f} 下{sec.bottom_margin.cm:.1f}cm')
    if rules.get('require_header') and not any(p.text.strip() for p in sec.header.paragraphs):
        issues.append('缺少页眉内容')
    if rules.get('require_footer') and not any(p.text.strip() for p in sec.footer.paragraphs):
        issues.append('缺少页脚内容')
    notes.append(f'页眉：{"有" if any(p.text.strip() for p in sec.header.paragraphs) else "无"}；页脚：{"有" if any(p.text.strip() for p in sec.footer.paragraphs) else "无"}')

    # 标题
    heads = defaultdict(list)
    for p in groups.get('heading', []):
        heads[p.style.name].append(p)
    for sname, plist in sorted(heads.items()):
        fonts = set()
        sizes = set()
        not_bold = 0
        for p in plist:
            for f, sz in run_fonts(p):
                if f: fonts.add(f)
                if sz: sizes.add(sz)
            # 加粗判断：run 级或段落样式级（样式定义加粗时不应误报）
            run_bold = any(r.bold for r in p.runs if r.text.strip())
            style_bold = bool(p.style.font.bold)
            if not (run_bold or style_bold):
                not_bold += 1
        allowed = rules['heading']['fonts']
        bad_fonts = [f for f in fonts if f not in allowed]
        if bad_fonts:
            issues.append(f'{sname}（{len(plist)} 个）字体不在规范内：{", ".join(bad_fonts)}（建议 {"/".join(allowed[:2])}）')
        if rules['heading'].get('bold_required') and not_bold:
            issues.append(f'{sname}（{len(plist)} 个）有 {not_bold} 个未加粗')
        notes.append(f'{sname}：{len(plist)} 个；字体 {"/".join(sorted(fonts)) or "继承"}；字号 {"/".join(sorted(str(s) for s in sizes)) or "继承"}')

    # 正文
    body = groups.get('body', [])
    if body:
        fonts = Counter(); sizes = Counter(); indents = Counter(); spacing = Counter()
        for p in body:
            for f, sz in run_fonts(p):
                fonts[f or '继承'] += 1
                if sz: sizes[sz] += 1
            ind = p.paragraph_format.first_line_indent
            indents[round(ind.pt) if ind else '无'] += 1
            ls = p.paragraph_format.line_spacing
            if isinstance(ls, float):
                spacing[ls] += 1
            elif hasattr(ls, 'pt'):
                spacing[f'{ls.pt}pt'] += 1
            else:
                spacing['继承'] += 1
        allowed_fonts = rules['body']['fonts']
        bad = [f for f in fonts if f not in allowed_fonts and f != '继承']
        if bad:
            issues.append(f'正文字体不规范：{", ".join(bad)}（规范 {"/".join(allowed_fonts[:2])}）')
        bad_sz = [s for s in sizes if s not in rules['body']['size_pt']]
        if bad_sz:
            issues.append(f'正文字号不规范：{", ".join(str(s) for s in bad_sz)}pt（规范 {rules["body"]["size_pt"]}pt）')
        if rules['body'].get('indent_required'):
            no_indent = indents.get('无', 0)
            if no_indent:
                issues.append(f'正文有 {no_indent}/{len(body)} 段缺少首行缩进（规范要求缩进）')
        notes.append(f'正文 {len(body)} 段；字体 {dict(fonts)}；字号 {dict(sizes)}')
        notes.append(f'首行缩进分布：{dict(indents)}（{"硬性要求" if rules["body"].get("indent_required") else "非硬性，仅统计"}）')
        notes.append(f'行距分布：{dict(spacing)}')

    # 列表 / 题注 / 代码（只统计，不判错）
    for g in ('list', 'caption', 'code', 'other'):
        plist = groups.get(g, [])
        if plist:
            notes.append(f'{g} 段落：{len(plist)} 个')

    # 表格
    tables = d.tables
    if tables:
        no_border = [i + 1 for i, t in enumerate(tables) if rules['table'].get('border_required') and not table_has_borders(t)]
        if no_border:
            issues.append(f'表格缺少边框：第 {no_border} 张')
        notes.append(f'表格 {len(tables)} 张；带边框 {len(tables) - len(no_border)}/{len(tables)}')

    # 图片
    if d.inline_shapes:
        avail = (sec.page_width.cm - sec.left_margin.cm - sec.right_margin.cm) - rules['image']['max_width_margin_cm']
        over = []
        max_w = 0
        for i, shp in enumerate(d.inline_shapes):
            w_cm = shp.width.cm
            max_w = max(max_w, w_cm)
            if w_cm > avail:
                over.append(i + 1)
        if over:
            issues.append(f'图片超出可用宽度（{avail:.1f}cm）：第 {over} 张')
        notes.append(f'图片 {len(d.inline_shapes)} 张；最大宽度 {max_w:.1f}cm（可用 {avail:.1f}cm）')

    return {'file': path, 'issues': issues, 'notes': notes, 'groups': {k: len(v) for k, v in groups.items()}}

def report(r):
    print(f'===== docx 格式规范校验 =====')
    print(f'文件：{r["file"]}')
    print(f'段落分组：{r["groups"]}')
    print()
    if r['issues']:
        print(f'▶ 与规范不符（{len(r["issues"])} 项）：')
        for i in r['issues']:
            print(f'    ⚠️ {i}')
    else:
        print('▶ 与规范不符：无 ✅')
    print()
    print('▶ 统计信息：')
    for n in r['notes']:
        print(f'    {n}')
    print()
    print('结论：' + ('格式规范 ✅' if not r['issues'] else f'{len(r["issues"])} 项需确认（缩进/行距等非硬性项仅作统计，不必强改）'))

if __name__ == '__main__':
    ap = argparse.ArgumentParser(description='docx 格式规范校验')
    ap.add_argument('docx', help='待校验的 .docx')
    ap.add_argument('--profile', default='technical', choices=['technical', 'academic'])
    ap.add_argument('--config', default=None, help='自定义规范 JSON（覆盖默认项）')
    ap.add_argument('--json', action='store_true')
    args = ap.parse_args()
    if not os.path.exists(args.docx):
        print(f'文件不存在：{args.docx}')
        sys.exit(1)
    rules = load_rules(args.profile, args.config)
    r = check(args.docx, rules)
    if args.json:
        print(json.dumps(r, ensure_ascii=False, indent=1))
    else:
        report(r)
    sys.exit(0 if not r['issues'] else 2)
